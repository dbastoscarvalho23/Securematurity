"""Registos em Excel por referencial (avaliações, Declarações de Aplicabilidade, listas legais).

Lê os dados da Biblioteca Legal (src/data/legal/*.json) e escreve docs/modelos/<referencial>/*.xlsx.
Uso: python3 registos.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[3]
LEGAL = ROOT / 'src' / 'data' / 'legal'
OUT = ROOT / 'docs' / 'modelos'
VERSION = 'Versão 1.0 · outubro de 2026'

NAVY, GOLD, TINT, INPUT = '070A13', 'CFA04A', 'F4F1EA', 'FFF8E1'
FONT = 'Arial'
thin = Side(style='thin', color='D9DEE7')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

ESTADO = ['Não iniciado', 'Em curso', 'Implementado', 'Não aplicável']


def load(name):
    return json.loads((LEGAL / f'{name}.json').read_text(encoding='utf-8'))


def readme(wb, title, lines, sources):
    ws = wb.active
    ws.title = 'Leia-me'
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws['A2'] = f'AnkoraOne · {VERSION}'
    ws['A2'].font = Font(name=FONT, size=10, color='5B6474')
    r = 4
    ws.cell(r, 1, 'Como preencher').font = Font(name=FONT, bold=True, color=NAVY)
    for line in lines:
        r += 1
        c = ws.cell(r, 1, f'• {line}')
        c.font = Font(name=FONT, size=10)
        c.alignment = WRAP
    r += 1
    c = ws.cell(r, 1, '• Células com fundo amarelo-claro: a preencher pela entidade. As restantes vêm da Biblioteca Legal e não devem ser alteradas.')
    c.font = Font(name=FONT, size=10)
    c.fill = PatternFill('solid', fgColor=INPUT)
    c.alignment = WRAP
    r += 2
    ws.cell(r, 1, 'Fontes').font = Font(name=FONT, bold=True, color=NAVY)
    for s in sources:
        r += 1
        c = ws.cell(r, 1, s)
        c.font = Font(name=FONT, size=10)
        c.alignment = WRAP
    ws.column_dimensions['A'].width = 120
    return ws


def sheet(wb, name, headers, widths, rows, inputs=(), estado_col=None, note=None):
    """Folha com cabeçalho, filtros, painel fixo e colunas de preenchimento destacadas."""
    ws = wb.create_sheet(name)
    start = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=FONT, size=10, italic=True, color='8A6620')
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.row_dimensions[1].height = 30
        ws.cell(1, 1).alignment = WRAP
        start = 2
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(start, j, h)
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, row in enumerate(rows, start + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONT, size=9)
            c.alignment = WRAP
            c.border = BORDER
            if j in inputs:
                c.fill = PatternFill('solid', fgColor=INPUT)
    last = start + len(rows)
    ws.freeze_panes = ws.cell(start + 1, 1)
    ws.auto_filter.ref = f'A{start}:{get_column_letter(len(headers))}{max(last, start + 1)}'
    if estado_col:
        dv = DataValidation(type='list', formula1='"' + ','.join(ESTADO) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(estado_col)
        dv.add(f'{col}{start + 1}:{col}{max(last, start + 200)}')
    return ws


def blank_rows(n, width, inputs):
    return [['' for _ in range(width)] for _ in range(n)]


def risks_sheet(wb, controls_label):
    headers = ['N.º', 'Risco (ameaça sobre o ativo)', 'Ativos', 'Probabilidade (1 a 5)', 'Impacto (1 a 5)', 'Nível (P × I)', 'Classificação', 'Tratamento', controls_label, 'Responsável', 'Prazo', 'Risco residual aceite por']
    widths = [6, 40, 22, 12, 12, 10, 14, 16, 26, 18, 12, 22]
    ws = sheet(wb, 'Riscos', headers, widths, [], inputs=(), note='Nível = Probabilidade × Impacto. Classificação: 1 a 6 baixo, 8 a 12 médio, 15 a 25 elevado (critério da AnkoraOne; ajuste aos critérios de aceitação aprovados pela entidade). A linha 3 é um exemplo.')
    example = [1, 'Exemplo: cifra por ransomware do servidor de ficheiros', 'Servidor de ficheiros', 3, 5, None, None, 'Reduzir', 'Cópias de segurança imutáveis; MFA', 'Responsável de TI', '2026-12-31', 'Responsável de cibersegurança']
    dv = DataValidation(type='list', formula1='"Reduzir,Evitar,Transferir,Aceitar"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('H3:H200')
    sc = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True)
    ws.add_data_validation(sc)
    sc.add('D3:E200')
    for r in range(3, 53):
        for j in range(1, 13):
            c = ws.cell(r, j, example[j - 1] if r == 3 else None)
            c.font = Font(name=FONT, size=9, italic=(r == 3), color=('5B6474' if r == 3 else '000000'))
            c.alignment = WRAP
            c.border = BORDER
            if j not in (6, 7):
                c.fill = PatternFill('solid', fgColor=INPUT)
        ws.cell(r, 6, f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),D{r}*E{r},"")')
        ws.cell(r, 7, f'=IF(F{r}="","",IF(F{r}>=15,"Elevado",IF(F{r}>=8,"Médio","Baixo")))')
        if r > 3:
            ws.cell(r, 1, r - 2)
    return ws


# ── 1. RJC e QNRCS ────────────────────────────────────────────────────────────
def rjc():
    q = load('qnrcs')
    wb = Workbook()
    readme(wb, 'Registos RJC e QNRCS', [
        'Avaliação QNRCS: uma linha por medida mínima e nível de conformidade (Reg. 756/2026, Anexo III). Filtre pela coluna «Nível». As entidades dos níveis Substancial e Elevado aplicam também as medidas dos níveis inferiores (Reg. 756/2026, art. 30.º, n.º 2).',
        'Ativos Internet: lista do Reg. 756/2026, art. 32.º, n.º 1. É informação classificada de grau reservado: guarde e partilhe com as regras correspondentes. Não a envie a ferramentas de IA.',
        'Estatística trimestral e incidentes significativos: alimentam o relatório anual (RJC, art. 30.º, n.º 1, alíneas b) e c)).',
        'Riscos: registo de apoio ao relatório de análise de riscos (RJC, art. 29.º; Reg. 756/2026, art. 31.º).',
    ], [q['source']['title'] + ' — ' + q['source']['ref'], q['source']['url'], 'Decreto-Lei n.º 125/2025 (RJC) — Diário da República, 1.ª série'])
    rows = []
    for o in q['objectives']:
        for cat in o['categories']:
            for c in cat['controls']:
                for lv, name in (('basico', 'Básico'), ('substancial', 'Substancial'), ('elevado', 'Elevado')):
                    m = (c.get('measures') or {}).get(lv)
                    if m:
                        rows.append([name, f"{o['code']} · {o['name']}", f"{cat['id']} · {cat['name']}", c['id'], c['statement'], m['measure'], m.get('verification', ''), '', '', '', ''])
    sheet(wb, 'Avaliação QNRCS', ['Nível', 'Objetivo', 'Categoria', 'Controlo', 'Declaração do controlo (Anexo I)', 'Medida mínima (Anexo III)', 'Critério de verificação', 'Estado', 'Evidência', 'Responsável', 'Prazo'],
          [11, 16, 22, 10, 40, 50, 40, 14, 30, 16, 12], rows, inputs=(8, 9, 10, 11), estado_col=8)
    sheet(wb, 'Ativos Internet (art. 32.º)', ['Serviço suportado', 'Nome do equipamento / software', 'Modelo / versão', 'Sistema operativo', 'Software do servidor aplicacional', 'Endereço IP (se aplicável)', 'FQDN (se aplicável)', 'Fabricante', 'Dependências de outros ativos'],
          [24, 26, 16, 18, 22, 18, 26, 16, 34], blank_rows(40, 9, None), inputs=tuple(range(1, 10)),
          note='RESERVADO — Reg. 756/2026, art. 32.º, n.º 1, alíneas a) a i). Versão inicial: até 31 de janeiro do ano seguinte à notificação da qualificação ou em 6 meses, o que vencer primeiro (n.º 3, alínea a)).')
    quarters = []
    for qt in ('1.º trimestre', '2.º trimestre', '3.º trimestre', '4.º trimestre'):
        for _ in range(5):
            quarters.append([qt, '', ''])
    ws = sheet(wb, 'Estatística trimestral', ['Trimestre', 'Tipo de incidente (taxonomia do CNCS)', 'Número'], [18, 46, 12], quarters, inputs=(2, 3),
               note='RJC, art. 30.º, n.º 1, alínea b): estatística trimestral de todos os incidentes, com o número e o tipo.')
    end = 2 + len(quarters)
    ws.cell(end + 2, 2, 'Total do ano').font = Font(name=FONT, bold=True, size=9)
    ws.cell(end + 2, 3, f'=SUM(C3:C{end})').font = Font(name=FONT, bold=True, size=9)
    sheet(wb, 'Incidentes significativos', ['Referência do incidente', 'Data de início', 'Data do fim do impacto', 'Utilizadores afetados', 'Duração (horas)', 'Distribuição geográfica', 'Impacto transfronteiriço (sim/não)', 'Problemas identificados e medidas'],
          [18, 14, 16, 14, 12, 26, 16, 44], blank_rows(20, 8, None), inputs=tuple(range(1, 9)),
          note='RJC, art. 30.º, n.º 1, alíneas c) e e): análise agregada dos incidentes com impacto significativo.')
    risks_sheet(wb, 'Medidas (QNRCS)')
    return wb


# ── 2. NIS2 e Reg. 2024/2690 ──────────────────────────────────────────────────
def nis2():
    x = load('regexec2690')
    wb = Workbook()
    readme(wb, 'Registos NIS2 e Reg. Exec. 2024/2690', [
        'Requisitos 2024/2690: os 161 requisitos técnicos e metodológicos do anexo, com a redação oficial. Diretamente aplicáveis às entidades do art. 1.º do regulamento; para as restantes servem de boa prática ao lado do QNRCS.',
        'Use a coluna «Estado» e junte a evidência. O resumo por alínea do art. 21.º, n.º 2, está no documento Word deste referencial (N3).',
    ], [f"{x['instrument'].get('title', 'Regulamento de Execução (UE) 2024/2690')}" if isinstance(x['instrument'], dict) else str(x['instrument']), 'EUR-Lex: https://eur-lex.europa.eu/legal-content/PT/TXT/?uri=CELEX%3A32024R2690'])
    sec = {s['id']: s for s in x['annex_sections']}
    rows = []
    for r in x['annex_requirements']:
        s = sec.get(r['id'].split('.')[0], {})
        rows.append([r['id'], r.get('section', ''), r.get('subsection', ''), ', '.join(s.get('nis2_letters', [])), r['text'], '', '', '', ''])
    sheet(wb, 'Requisitos 2024-2690', ['Ponto', 'Secção', 'Subsecção', 'Art. 21.º, n.º 2, alínea', 'Requisito (redação oficial)', 'Estado', 'Evidência', 'Responsável', 'Prazo'],
          [8, 28, 30, 10, 70, 14, 30, 16, 12], rows, inputs=(6, 7, 8, 9), estado_col=6)
    return wb


# ── 3. DORA ───────────────────────────────────────────────────────────────────
def dora():
    d = load('dora')
    wb = Workbook()
    readme(wb, 'Registos DORA', [
        'Obrigações DORA: 111 obrigações do Regulamento (UE) 2022/2554 e dos atos da Biblioteca Legal, por pilar. A coluna «Texto» é a redação oficial.',
        'Registo de informação: organização da recolha por modelo do ITS 2024/2956. O envio segue o formato técnico das autoridades de supervisão.',
    ], [f'{s}' for s in d['meta']['sources']])
    PIL = {'governacao': 'Governação', 'gestao_risco_tic': 'Gestão do risco TIC', 'incidentes': 'Incidentes', 'testes': 'Testes de resiliência', 'terceiros_tic': 'Terceiros prestadores de TIC', 'partilha_informacao': 'Partilha de informação'}
    rows = [[PIL.get(o['pillar'], o['pillar']), o['ref'], o['title_pt'], o.get('summary_pt', ''), o['text'], '', '', ''] for o in d['obligations']]
    sheet(wb, 'Obrigações DORA', ['Pilar', 'Referência', 'Obrigação', 'Síntese', 'Texto (redação oficial)', 'Estado', 'Evidência', 'Responsável'],
          [20, 22, 34, 44, 70, 14, 30, 16], rows, inputs=(6, 7, 8), estado_col=6)
    rows = [[t['code'], t['name'], t.get('description', ''), t.get('instructions_ref', ''), '', '', ''] for t in d['its_2956_templates']]
    sheet(wb, 'Registo de informação', ['Modelo', 'Nome', 'O que identifica', 'Instruções', 'Fonte interna dos dados', 'Responsável', 'Estado'],
          [10, 34, 70, 22, 28, 16, 14], rows, inputs=(5, 6, 7), estado_col=7)
    return wb


# ── 4 e 5. ISO ────────────────────────────────────────────────────────────────
def iso(name, std, link_key, link_label):
    d = load(name)
    wb = Workbook()
    readme(wb, f'Registos {std}', [
        'Declaração de Aplicabilidade: um controlo por linha do Anexo A. Indique se é aplicável, a justificação (inclusão ou exclusão), se está implementado e o método.',
        'Títulos e descrições são redação da AnkoraOne; a norma não é reproduzida. Para certificação use o texto oficial.',
        'Riscos: registo de apoio à avaliação e ao tratamento de riscos (cláusulas 6.1.2, 6.1.3, 8.2 e 8.3).',
    ], [d['source']['title'], d['source']['edition'], d['source']['note']])
    controls = [c for c in d['items'] if c['kind'] == 'control']
    rows = [[f"A.{c['ref']}", c['group_name'], c['title'], c['description'], c.get('evidence', ''), ', '.join(c.get(link_key, []) or []), '', '', '', '', ''] for c in controls]
    ws = sheet(wb, 'Declaração de Aplicabilidade', ['Controlo', 'Tema', 'Título (guia AnkoraOne)', 'Descrição (guia AnkoraOne)', 'Evidência sugerida', link_label, 'Aplicável (sim/não)', 'Justificação', 'Implementado (sim/parcial/não)', 'Método de implementação', 'Documento / evidência'],
               [10, 22, 30, 50, 34, 18, 12, 36, 14, 30, 28], rows, inputs=(7, 8, 9, 10, 11))
    for col, opts in ((7, 'Sim,Não'), (9, 'Sim,Parcial,Não')):
        dv = DataValidation(type='list', formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        L = get_column_letter(col)
        dv.add(f'{L}2:{L}{len(rows) + 1}')
    risks_sheet(wb, 'Controlos (Anexo A ou próprios)')
    docs = [[c['ref'], c['title'], '', '', ''] for c in d['items'] if c['kind'] == 'clause' and c.get('documented')]
    sheet(wb, 'Informação documentada', ['Cláusula', 'Tema', 'Documento da entidade', 'Local', 'Versão e data'], [10, 44, 34, 30, 16], docs, inputs=(3, 4, 5))
    return wb, d


def iso27001():
    wb, _ = iso('iso27001', 'ISO/IEC 27001:2022', 'qnrcs', 'Controlos QNRCS relacionados')
    return wb


def iso42001():
    wb, d = iso('iso42001', 'ISO/IEC 42001:2023', 'iso27001', 'Ligação à ISO/IEC 27001')
    sheet(wb, 'Sistemas de IA', ['Sistema de IA', 'Finalidade', 'Papel da organização', 'Fornecedor / modelo', 'Dados utilizados', 'Pessoas afetadas', 'Nível de automação', 'Avaliação de impacto (data)', 'Responsável', 'Estado'],
          [24, 30, 18, 22, 26, 24, 16, 16, 18, 14], blank_rows(25, 10, None), inputs=tuple(range(1, 11)), estado_col=10,
          note='Inventário de sistemas de IA no âmbito (cláusula 4.3). Papel: fornecedor, produtor, cliente ou utilizador (4.1).')
    rows = [['Objetivo', o['id'], o['title'], o['description'], '', ''] for o in d['objectives']] + [['Fonte de risco', r['id'], r['title'], r['description'], '', ''] for r in d['risk_sources']]
    sheet(wb, 'Anexo C', ['Tipo', 'Ref.', 'Tema (guia AnkoraOne)', 'Descrição (guia AnkoraOne)', 'Relevante? (sim/não)', 'Como é tratado'], [14, 8, 30, 60, 14, 40], rows, inputs=(5, 6))
    return wb


# ── 6. TISAX ──────────────────────────────────────────────────────────────────
def tisax():
    i = load('isa6')
    inst = i['instrument']
    wb = Workbook()
    readme(wb, 'Autoavaliação TISAX — VDA ISA 6.0.3', [
        'Uma linha por controlo do catálogo ISA. As colunas com o texto do catálogo estão em inglês e sem alterações, como exige a licença.',
        'Preencha o nível de maturidade alvo e atual (0 a 5), a evidência e o responsável. Os níveis de maturidade estão na folha «Níveis de maturidade».',
        'Os requisitos «high» e «very high» só se aplicam aos objetivos de avaliação com esse nível de proteção.',
    ], [f"{inst['short']} — {inst.get('cover_version_line', '')}", inst['official_url'], f"Licença: {inst['license_name']}", 'Atribuição: ver a folha «Atribuição e licença».'])
    j = lambda v: '\n'.join(v) if isinstance(v, list) else (v or '')
    rows = [[c['id'], c['chapter'], c['subchapter'], c['question'], c.get('objective', ''), j(c.get('must')), j(c.get('should')), j(c.get('high')), j(c.get('very_high')), '', '', '', ''] for c in i['controls']]
    ws = sheet(wb, 'Autoavaliação ISA 6', ['ID', 'Chapter', 'Subchapter', 'Control question', 'Objective', 'Requirements (must)', 'Requirements (should)', 'Additional requirements for high protection needs', 'Additional requirements for very high protection needs', 'Maturidade alvo (0 a 5)', 'Maturidade atual (0 a 5)', 'Evidência', 'Responsável'],
               [8, 20, 24, 36, 44, 50, 50, 44, 44, 12, 12, 30, 16], rows, inputs=(10, 11, 12, 13))
    dv = DataValidation(type='whole', operator='between', formula1='0', formula2='5', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'J2:K{len(rows) + 1}')
    sheet(wb, 'Níveis de maturidade', ['Nível', 'Name', 'Description', 'Definition', 'Possible evidence'], [8, 16, 60, 60, 40],
          [[m['level'], m['name'], m.get('text', ''), m.get('definition', ''), m.get('possible_evidence') or ''] for m in i['maturity_levels']])
    lic = wb.create_sheet('Atribuição e licença')
    lic['A1'] = inst.get('attribution', '')
    lic['A3'] = inst.get('publisher_text', '')
    lic['A5'] = inst.get('license_text', '')
    for cell in ('A1', 'A3', 'A5'):
        lic[cell].alignment = WRAP
        lic[cell].font = Font(name=FONT, size=9)
    lic.column_dimensions['A'].width = 140
    return wb


PACKS = [
    ('rjc-qnrcs', 'AnkoraOne-Registos-RJC-QNRCS.xlsx', rjc),
    ('nis2', 'AnkoraOne-Registos-NIS2-Reg2024-2690.xlsx', nis2),
    ('dora', 'AnkoraOne-Registos-DORA.xlsx', dora),
    ('iso27001', 'AnkoraOne-Registos-ISO27001.xlsx', iso27001),
    ('iso42001', 'AnkoraOne-Registos-ISO42001.xlsx', iso42001),
    ('tisax', 'AnkoraOne-Autoavaliacao-TISAX-ISA6.xlsx', tisax),
]

if __name__ == '__main__':
    for folder, file, fn in PACKS:
        d = OUT / folder
        d.mkdir(parents=True, exist_ok=True)
        fn().save(d / file)
        print('→', f'{folder}/{file}')
